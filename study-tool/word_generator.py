"""
إنشاء ملفات Word
=================
يحول نتائج AI إلى 3 ملفات Word منسقة.
"""

import os
import re
from typing import Dict

from docx import Document
from docx.shared import Pt, Inches, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn


def _set_rtl_paragraph(paragraph):
    """تفعيل RTL للفقرة (للنصوص العربية)."""
    pPr = paragraph._element.get_or_add_pPr()
    bidi = pPr.makeelement(qn('w:bidi'), {})
    pPr.append(bidi)


def _create_base_document(title: str, subtitle: str = "") -> Document:
    """إنشاء مستند Word أساسي مع إعدادات مشتركة."""
    doc = Document()

    # إعداد الصفحة
    section = doc.sections[0]
    section.page_width = Cm(21)  # A4
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)

    # إعداد الخط الافتراضي
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Calibri'
    font.size = Pt(12)
    font.color.rgb = RGBColor(0x33, 0x33, 0x33)

    # العنوان الرئيسي
    heading = doc.add_heading(title, level=0)
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER

    if subtitle:
        sub = doc.add_paragraph(subtitle)
        sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
        sub.runs[0].font.size = Pt(11)
        sub.runs[0].font.color.rgb = RGBColor(0x66, 0x66, 0x66)

    doc.add_paragraph("")  # فراغ

    return doc


def _add_content_to_doc(doc: Document, content: str):
    """
    إضافة المحتوى إلى المستند مع تنسيق ذكي.
    يتعرف على العناوين (Slide headers) والنقاط.
    """
    lines = content.split('\n')

    for line in lines:
        stripped = line.strip()
        if not stripped:
            doc.add_paragraph("")  # سطر فارغ
            continue

        # التعرف على عناوين السلايدات
        slide_match = re.match(r'^(?:===\s*)?(?:Slide|سلايد)\s*(\d+)\s*(?:===)?:?\s*(.*)', stripped, re.IGNORECASE)
        if slide_match:
            slide_num = slide_match.group(1)
            slide_title = slide_match.group(2).strip(' =:')
            heading_text = f"Slide {slide_num}"
            if slide_title:
                heading_text += f": {slide_title}"
            h = doc.add_heading(heading_text, level=2)
            continue

        # عناوين بـ ### أو ** **
        if stripped.startswith('### ') or stripped.startswith('## '):
            text = stripped.lstrip('#').strip()
            doc.add_heading(text, level=3)
            continue

        # نقاط
        bullet_match = re.match(r'^[-•*]\s+(.*)', stripped)
        if bullet_match:
            p = doc.add_paragraph(bullet_match.group(1), style='List Bullet')
            continue

        # نقاط مرقمة
        numbered_match = re.match(r'^(\d+)[.)]\s+(.*)', stripped)
        if numbered_match:
            p = doc.add_paragraph(numbered_match.group(2), style='List Number')
            continue

        # خط عريض **text**
        if stripped.startswith('**') and stripped.endswith('**'):
            p = doc.add_paragraph()
            run = p.add_run(stripped.strip('*'))
            run.bold = True
            run.font.size = Pt(13)
            continue

        # نص عادي
        p = doc.add_paragraph(stripped)


def create_translation_doc(content: str, pdf_name: str, output_path: str):
    """
    إنشاء ملف الترجمة (الملف الأول).

    Args:
        content: نتيجة AI للترجمة
        pdf_name: اسم ملف PDF الأصلي
        output_path: مسار الملف الناتج
    """
    doc = _create_base_document(
        title="Translation Reference",
        subtitle=f"مرجع الترجمة - {pdf_name}"
    )

    _add_content_to_doc(doc, content)
    doc.save(output_path)


def create_simplified_doc(content: str, pdf_name: str, output_path: str):
    """
    إنشاء ملف التبسيط (الملف الثاني).

    Args:
        content: نتيجة AI للتبسيط
        pdf_name: اسم ملف PDF الأصلي
        output_path: مسار الملف الناتج
    """
    doc = _create_base_document(
        title="Simplified Study Guide",
        subtitle=f"دليل الدراسة المبسط - {pdf_name}"
    )

    _add_content_to_doc(doc, content)
    doc.save(output_path)


def create_flashcards_doc(content: str, pdf_name: str, output_path: str):
    """
    إنشاء ملف البطاقات التعليمية (الملف الثالث).

    Args:
        content: نتيجة AI للبطاقات
        pdf_name: اسم ملف PDF الأصلي
        output_path: مسار الملف الناتج
    """
    doc = _create_base_document(
        title="Flash Cards",
        subtitle=f"بطاقات المذاكرة - {pdf_name}"
    )

    # تنسيق خاص للبطاقات: إضافة خطوط فاصلة وتنسيق Front/Back
    lines = content.split('\n')
    card_count = 0

    for line in lines:
        stripped = line.strip()
        if not stripped:
            continue

        # بداية بطاقة جديدة
        card_header = re.match(r'^(?:Card|بطاقة)\s*#?\s*(\d+)', stripped, re.IGNORECASE)
        if card_header:
            card_count += 1
            if card_count > 1:
                # خط فاصل بين البطاقات
                doc.add_paragraph("─" * 50)
            h = doc.add_heading(f"Card #{card_header.group(1)}", level=3)
            continue

        # Front
        front_match = re.match(r'^Front\s*\(?\s*الوجه\s*\)?\s*:?\s*(.*)', stripped, re.IGNORECASE)
        if front_match:
            p = doc.add_paragraph()
            run = p.add_run("Front (الوجه): ")
            run.bold = True
            run.font.color.rgb = RGBColor(0x00, 0x70, 0xC0)
            if front_match.group(1):
                p.add_run(front_match.group(1))
            continue

        # Back
        back_match = re.match(r'^Back\s*\(?\s*الخلف\s*\)?\s*:?\s*(.*)', stripped, re.IGNORECASE)
        if back_match:
            p = doc.add_paragraph()
            run = p.add_run("Back (الخلف): ")
            run.bold = True
            run.font.color.rgb = RGBColor(0x00, 0xA0, 0x50)
            if back_match.group(1):
                p.add_run(back_match.group(1))
            continue

        # نص عادي (تكملة Front/Back أو ملاحظات)
        _add_content_to_doc(doc, stripped)

    doc.save(output_path)


def generate_all_documents(
    ai_results: Dict[str, str],
    pdf_name: str,
    output_folder: str,
    file_names: Dict[str, str],
) -> list:
    """
    إنشاء الملفات الثلاثة في المجلد المحدد.

    Args:
        ai_results: نتائج AI {"translation": ..., "simplified": ..., "flashcards": ...}
        pdf_name: اسم ملف PDF
        output_folder: مسار المجلد الناتج
        file_names: أسماء الملفات {"file1": ..., "file2": ..., "file3": ...}

    Returns:
        قائمة مسارات الملفات المُنشأة
    """
    os.makedirs(output_folder, exist_ok=True)

    created_files = []

    # الملف 1: الترجمة
    path1 = os.path.join(output_folder, file_names.get("file1", "01_Translation.docx"))
    create_translation_doc(ai_results["translation"], pdf_name, path1)
    created_files.append(path1)

    # الملف 2: التبسيط
    path2 = os.path.join(output_folder, file_names.get("file2", "02_Simplified.docx"))
    create_simplified_doc(ai_results["simplified"], pdf_name, path2)
    created_files.append(path2)

    # الملف 3: البطاقات
    path3 = os.path.join(output_folder, file_names.get("file3", "03_FlashCards.docx"))
    create_flashcards_doc(ai_results["flashcards"], pdf_name, path3)
    created_files.append(path3)

    return created_files
